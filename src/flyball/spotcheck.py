"""Side-by-side comparison of raw Excel cells and the cleaned heat row.

Used for the manual spot-check of Phase 2: the raw values are read straight
from the workbook with openpyxl, independently of the cleaning pipeline.
"""

import openpyxl
import pandas as pd

from flyball import config


def raw_row(sheet: str, excel_row: int) -> dict:
    """Raw cell values of columns A-T for one sheet row."""
    for path in config.COMPETITION_XLSX.values():
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        if sheet in wb.sheetnames:
            ws = wb[sheet]
            cells = next(ws.iter_rows(min_row=excel_row, max_row=excel_row,
                                      max_col=len(config.SHEET_COLUMNS), values_only=True))
            return dict(zip(config.SHEET_COLUMNS, cells))
    raise KeyError(sheet)


def compare(heats: pd.DataFrame, sheet: str, excel_row: int) -> pd.DataFrame:
    """Raw value vs cleaned value (and code) for every column of one row."""
    raw = raw_row(sheet, excel_row)
    row = heats[(heats["sheet"] == sheet) & (heats["excel_row"] == excel_row)].iloc[0]
    out = []
    for col in config.SHEET_COLUMNS[1:]:
        clean_value = row["heat_time"] if col == "when" else row.get(col)
        out.append({"column": col, "raw": raw[col], "clean": clean_value,
                    "code": row.get(f"{col}_code", "")})
    for col in ["lane", "heat_id", "race_id", "heat_in_race", "clean_time", "rerun_time",
                "n_faults", "is_practice", "row_realigned"]:
        out.append({"column": col, "raw": "", "clean": row[col], "code": ""})
    return pd.DataFrame(out)
